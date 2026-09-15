# -*- coding: utf-8 -*-
"""沿用项目 openpyxl 工具生成真实 XLSX，全量数据由报表服务提供。"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from .task_type_report import ReportError

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
MAX_EXPORT_ROWS = 50000
COLUMNS = (
    ("first_bbk_id", "分行号"),
    ("first_bbk_name", "分行名称"),
    ("org_id", "网点号"),
    ("org_name", "网点名称"),
    ("user_id", "客户经理ID"),
    ("user_name", "客户经理姓名"),
    ("sapid", "SAP号"),
    ("pst_lvl", "SAP岗位"),
    ("skill_id", "技能ID"),
    ("cn_name", "技能名称"),
    ("task_type_name", "任务类型"),
    ("skill_count", "技能数"),
    ("permission_manager_count", "有权限客户经理数"),
    ("active_manager_count", "活跃客户经理数"),
    ("suc_execute_job", "成功任务数"),
    ("read_tasks", "已查看任务数"),
    ("read_rate", "任务查看率（%）"),
    ("recommended_customers", "方案客户数"),
    ("read_customer_count", "已查看方案客户数"),
    ("plan_read_rate", "方案查看率（%）"),
    ("insight_customer_count", "洞察客户数"),
    ("click_to_insight_rate", "洞察覆盖率（%）"),
    ("insight_count", "点击客户洞察总次数"),
    ("phone_customer_count", "电访客户数"),
    ("click_to_phone_rate", "电访覆盖率（%）"),
    ("phone_count", "点击去电访总次数"),
)


def _cell(sheet, value, percentage=False):
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    elif value is not None:
        if percentage:
            cell.value = value / 100
            cell.number_format = "0.00%"
        else:
            cell.number_format = "#,##0"
    return cell


def export_task_type_report(report, params, bbk_id: str) -> bytes:
    if len(report.items) > MAX_EXPORT_ROWS:
        raise ReportError(
            413,
            "report_export_too_large",
            "导出超过50000行，请缩小时间或机构范围。",
        )
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(
        "技能明细" if report.skill_detail else "统计报表"
    )
    sheet.freeze_panes = "A2"
    headers = [_cell(sheet, label) for _, label in COLUMNS]
    for cell in headers:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    sheet.append(headers)
    for row in report.items:
        sheet.append(
            [
                _cell(sheet, getattr(row, field), field.endswith("rate"))
                for field, _ in COLUMNS
            ]
        )
    metadata = workbook.create_sheet("筛选与口径")
    records = [
        ("来源", report.source_id),
        ("访问分行范围", bbk_id),
        ("名单快照", report.sync_date),
        ("实际分行筛选", report.resolved_filters.first_bbk_id),
        ("实际网点筛选", report.resolved_filters.org_id),
        ("全量行数", report.total),
        ("口径", report.metric_version),
        (
            "比例",
            "百分点按百分比格式展示，零分母留空；期间比率不是同批次转化率。",
        ),
        ("主动任务查看数", "等于成功任务数，不代表埋点观测。"),
        ("技能明细", "多技能任务分别计入关联技能，技能明细不可相加作为汇总。"),
    ]
    records.extend(
        (name, str(value) if value is not None else None)
        for name, value in params.model_dump(mode="json").items()
    )
    for key, value in records:
        metadata.append([_cell(metadata, key), _cell(metadata, value)])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
