# -*- coding: utf-8 -*-
"""Verify exported branch data, formatting, numeric sorting and formulas."""

from io import BytesIO

from openpyxl import load_workbook
import pytest

from monitor.app.models.cron import CronBranchRankingItem
from monitor.app.services.cron.branch_export import (
    branch_values,
    export_branch_dimension,
)


def item(name="分行", **values):
    return CronBranchRankingItem(bbk_id="100", bbk_name=name, **values)


def workbook_rows(items, **sort):
    return load_workbook(
        BytesIO(export_branch_dimension(items, **sort)),
    ).active


def test_empty_workbook_has_two_header_rows_and_all_columns():
    sheet = workbook_rows([])
    assert (sheet.max_row, sheet.max_column) == (2, 23)
    assert {str(r) for r in sheet.merged_cells.ranges} == {
        "A1:A2",
        "B1:B2",
        "C1:G1",
        "H1:P1",
        "Q1:W1",
    }
    assert sheet["A1"].value is None
    assert sheet["B1"].value == "分行名称"
    assert sheet["C1"].value == "任务信息"
    assert sheet["H1"].value == "by客户经理"
    assert sheet["Q1"].value == "by客户"
    assert sheet["G2"].value == "任务已读率"
    assert sheet["W2"].value == "接触客户率"


def test_metrics_and_formats_follow_console():
    sheet = workbook_rows(
        [
            item(
                skill_count=1234,
                total_tasks=2345,
                success_count=3456,
                read_tasks=1728,
                involved_managers=20,
                result_view_managers=10,
                plan_managers=4,
                insight_managers=2,
                phone_managers=1,
                recommended_customers=8,
                viewed_customers=3,
                insight_customers=2,
                phone_customers=1,
                contacted_customers=2,
                contact_rate=0.123456,
            ),
        ],
    )
    assert [c.value for c in sheet[3]] == [
        1,
        "分行",
        1234,
        2345,
        3456,
        1728,
        0.5,
        20,
        10,
        0.5,
        4,
        0.4,
        2,
        0.2,
        1,
        0.1,
        8,
        3,
        0.375,
        2,
        1,
        2,
        0.1235,
    ]
    assert sheet["C3"].number_format == "#,##0"
    assert all(sheet[f"{c}3"].number_format == "0.00%" for c in "GJLNPSW")


def test_percentage_values_are_plain_numbers():
    values = branch_values(
        item(
            success_count=4,
            read_tasks=1,
            involved_managers=4,
            result_view_managers=2,
            plan_managers=1,
            insight_managers=1,
            phone_managers=1,
            recommended_customers=4,
            viewed_customers=1,
            contact_rate=0.125,
        ),
    )
    percentage_indexes = (4, 7, 9, 11, 13, 16, 20)
    assert all(
        isinstance(values[index], (int, float)) for index in percentage_indexes
    )


@pytest.mark.parametrize(
    "sort_order,expected",
    [
        ("asc", ["two", "ten", "ten-again"]),
        ("desc", ["ten", "ten-again", "two"]),
    ],
)
def test_numeric_sort_is_stable_and_reranks(sort_order, expected):
    sheet = workbook_rows(
        [
            item("ten", total_tasks=10),
            item("two", total_tasks=2),
            item("ten-again", total_tasks=10),
        ],
        sort_by="totalTasks",
        sort_order=sort_order,
    )
    assert [sheet.cell(r, 2).value for r in range(3, 6)] == expected
    assert [sheet.cell(r, 1).value for r in range(3, 6)] == [1, 2, 3]


def test_percentage_sort_uses_displayed_precision():
    sheet = workbook_rows(
        [
            item("first", contact_rate=0.123444),
            item("second", contact_rate=0.123441),
        ],
        sort_by="contactRate",
        sort_order="asc",
    )
    assert sheet["B3"].value == "first"
    assert sheet["W3"].value == sheet["W4"].value == 0.1234


def test_zero_divisors_and_untrusted_branch_names():
    sheet = workbook_rows(
        [item('=HYPERLINK("https://example.org")'), item("")],
    )
    assert sheet["B3"].data_type == "s"
    assert sheet["B3"].value.startswith("=HYPERLINK")
    assert sheet["B4"].value == "100"
    assert all(sheet[f"{c}3"].value == 0 for c in "GJLNPSW")


def test_export_does_not_truncate_rows_or_change_default_order():
    sheet = workbook_rows([item(str(i)) for i in range(121)])
    assert sheet.max_row == 123
    assert sheet["B3"].value == "0"
    assert sheet["B123"].value == "120"
